"""Seed a deterministic 20-material demo dataset for manual/API verification.

The dataset intentionally covers every supported upload format:
- 10 document files: 4 PDF, 3 DOCX, 3 XLSX
- 10 supporting files: 3 TXT, 3 MARKDOWN, 2 CSV, 2 PPTX

Run against a running backend:
    venv/bin/python scripts/seed_demo_dataset.py

The script is idempotent: it reuses the demo subject and skips files whose
filename already exists for that subject.
"""
from __future__ import annotations

import csv
import io
import json
import sys
import textwrap
import time
from pathlib import Path

import docx
import fitz
import openpyxl
import requests
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "storage" / "demo_seed"
API = "http://127.0.0.1:8000/api"
SUBJECT_NAME = "Demo 20 Materi - Biologi Sel"
SUBJECT_META = {
    "semester": "1",
    "lecturer": "Tim Demo Kabayan",
    "description": (
        "Dataset uji end-to-end berisi 20 materi biologi sel dan genetika. "
        "Sepuluh file dokumen digunakan untuk menguji PDF, DOCX, dan XLSX; "
        "sepuluh file pendukung mencakup TXT, Markdown, CSV, dan PPTX."
    ),
}


TOPICS = [
    {
        "title": "Struktur dan Fungsi Membran Sel",
        "summary": "Membran sel adalah lapisan tipis yang membatasi sel dan mengatur pertukaran zat. Model mosaik fluida menjelaskan bahwa membran tersusun dari fosfolipid, protein, kolesterol, dan karbohidrat yang dapat bergerak lateral.",
        "points": [
            "Fosfolipid membentuk bilayer dengan ekor hidrofobik menghadap ke dalam dan kepala hidrofilik menghadap ke luar.",
            "Protein integral menembus bilayer, sedangkan protein perifer menempel pada permukaan membran.",
            "Kolesterol menjaga fluiditas membran pada suhu yang berubah-ubah.",
            "Glikoprotein dan glikolipid berperan dalam pengenalan sel dan interaksi antar sel.",
            "Membran bersifat selektif permeabel sehingga tidak semua molekul dapat melewatinya.",
        ],
        "terms": [
            ("Bilayer", "Dua lapis fosfolipid yang menyusun kerangka membran."),
            ("Mosaik fluida", "Model yang menggambarkan membran sebagai susunan komponen yang dinamis."),
            ("Selektif permeabel", "Sifat membran yang hanya meloloskan molekul tertentu."),
        ],
        "case": "Pada suhu rendah, kolesterol mencegah asam lemak saling berdekatan terlalu rapat sehingga membran tidak mengeras. Pada suhu tinggi, kolesterol membatasi gerakan fosfolipid agar membran tidak terlalu cair.",
    },
    {
        "title": "Transport Pasif dan Transport Aktif",
        "summary": "Transport pasif memindahkan zat tanpa energi ATP, sedangkan transport aktif membutuhkan energi untuk melawan gradien konsentrasi. Difusi, osmosis, dan difusi terfasilitasi termasuk transport pasif.",
        "points": [
            "Difusi sederhana terjadi pada molekul kecil nonpolar seperti oksigen dan karbon dioksida.",
            "Osmosis adalah perpindahan air melalui membran selektif permeabel dari konsentrasi zat terlarut rendah ke tinggi.",
            "Difusi terfasilitasi menggunakan protein kanal atau protein pembawa tanpa ATP.",
            "Pompa natrium-kalium adalah contoh transport aktif yang memakai ATP.",
            "Endositosis dan eksositosis memindahkan materi dalam vesikel.",
        ],
        "terms": [
            ("Gradien konsentrasi", "Perbedaan konsentrasi zat antara dua daerah."),
            ("Pompa natrium-kalium", "Protein membran yang memompa Na+ keluar dan K+ masuk dengan ATP."),
            ("Osmosis", "Perpindahan air melewati membran semipermeabel."),
        ],
        "case": "Sel darah merah yang ditempatkan dalam larutan hipotonik akan menyerap air dan dapat mengalami hemolisis. Dalam larutan hipertonik, sel kehilangan air dan mengerut.",
    },
    {
        "title": "Organel Sel dan Fungsinya",
        "summary": "Sel eukariot memiliki organel dengan fungsi khusus. Nukleus menyimpan DNA, ribosom menyintesis protein, retikulum endoplasma memproses protein dan lipid, serta aparatus Golgi mengemas dan mengirim produk sel.",
        "points": [
            "Mitokondria menghasilkan ATP melalui respirasi sel dan memiliki DNA sendiri.",
            "Kloroplas ditemukan pada sel tumbuhan dan berperan dalam fotosintesis.",
            "Lisosom mengandung enzim hidrolitik untuk mencerna makromolekul.",
            "Vakuola pada sel tumbuhan menyimpan air, ion, dan zat lain serta menjaga turgor.",
            "Sitoskeleton memberi bentuk sel dan membantu transport intraseluler.",
        ],
        "terms": [
            ("Nukleus", "Organel yang menyimpan materi genetik dan mengatur aktivitas sel."),
            ("Ribosom", "Tempat sintesis protein."),
            ("Lisosom", "Organel pencernaan intraseluler."),
        ],
        "case": "Sel pankreas memiliki retikulum endoplasma kasar dan aparatus Golgi yang berkembang baik karena aktif menyintesis dan mensekresikan enzim pencernaan.",
    },
    {
        "title": "Mitokondria dan Respirasi Sel",
        "summary": "Mitokondria adalah tempat utama respirasi aerob yang menghasilkan ATP. Glikolisis terjadi di sitosol, sedangkan siklus Krebs dan fosforilasi oksidatif berlangsung di mitokondria.",
        "points": [
            "Glikolisis memecah glukosa menjadi piruvat dan menghasilkan sedikit ATP serta NADH.",
            "Siklus Krebs menghasilkan NADH, FADH2, dan karbon dioksida.",
            "Rantai transpor elektron menggunakan NADH dan FADH2 untuk memompa proton.",
            "Gradien proton menggerakkan ATP sintase untuk menghasilkan ATP.",
            "Oksigen berfungsi sebagai akseptor elektron terakhir pada respirasi aerob.",
        ],
        "terms": [
            ("ATP", "Molekul pembawa energi utama sel."),
            ("Siklus Krebs", "Rangkaian reaksi di matriks mitokondria yang menghasilkan elektron pembawa energi."),
            ("Fosforilasi oksidatif", "Pembentukan ATP yang digerakkan oleh transpor elektron."),
        ],
        "case": "Jika pasokan oksigen berkurang, sel dapat beralih ke fermentasi laktat. Fermentasi menghasilkan ATP lebih sedikit dibandingkan respirasi aerob, tetapi memungkinkan glikolisis terus berjalan.",
    },
    {
        "title": "Kloroplas dan Fotosintesis",
        "summary": "Fotosintesis mengubah energi cahaya menjadi energi kimia pada kloroplas. Reaksi terang menghasilkan ATP dan NADPH, sedangkan siklus Calvin menggunakan keduanya untuk mengikat karbon dioksida.",
        "points": [
            "Klorofil menyerap cahaya terutama pada spektrum merah dan biru.",
            "Reaksi terang berlangsung di membran tilakoid dan melibatkan fotosistem II dan I.",
            "Fotolisis air menghasilkan oksigen, proton, dan elektron.",
            "Siklus Calvin berlangsung di stroma dan menghasilkan glukosa.",
            "Faktor yang memengaruhi fotosintesis meliputi intensitas cahaya, suhu, dan kadar CO2.",
        ],
        "terms": [
            ("Tilakoid", "Membran tempat reaksi terang fotosintesis."),
            ("Stroma", "Cairan di dalam kloroplas tempat siklus Calvin."),
            ("NADPH", "Pembawa elektron dan hidrogen untuk sintesis gula."),
        ],
        "case": "Tanaman yang kekurangan magnesium akan mengalami klorosis karena magnesium merupakan bagian penting dari molekul klorofil.",
    },
    {
        "title": "Siklus Sel dan Regulasi",
        "summary": "Siklus sel terdiri atas interfase dan fase pembelahan. Interfase mencakup G1, S, dan G2, sedangkan fase M mencakup mitosis dan sitokinesis.",
        "points": [
            "Fase G1 adalah pertumbuhan sel dan persiapan sintesis DNA.",
            "Fase S adalah replikasi DNA.",
            "Fase G2 adalah persiapan pembelahan dan pemeriksaan kerusakan DNA.",
            "Checkpoint G1/S memastikan DNA layak direplikasi.",
            "Siklus sel diatur oleh siklin, kinase dependen siklin, dan protein penghambat.",
        ],
        "terms": [
            ("Interfase", "Fase terpanjang siklus sel saat sel tumbuh dan mereplikasi DNA."),
            ("Checkpoint", "Titik pemeriksaan yang memastikan kondisi sel aman untuk melanjutkan siklus."),
            ("Siklin", "Protein yang mengatur aktivitas CDK pada fase tertentu."),
        ],
        "case": "Kerusakan DNA yang tidak dapat diperbaiki dapat memicu apoptosis melalui jalur p53. Hilangnya fungsi p53 berkaitan dengan pertumbuhan sel yang tidak terkendali.",
    },
    {
        "title": "Mitosis dan Meiosis",
        "summary": "Mitosis menghasilkan dua sel anak yang identik dengan sel induk, sedangkan meiosis menghasilkan empat sel haploid yang berbeda secara genetik. Mitosis berperan dalam pertumbuhan dan perbaikan jaringan.",
        "points": [
            "Profase ditandai oleh kondensasi kromosom dan hilangnya membran nukleus.",
            "Metafase menyusun kromosom di bidang ekuator sel.",
            "Anafase memisahkan kromatid saudara menuju kutub berlawanan.",
            "Telofase diikuti sitokinesis yang membagi sitoplasma.",
            "Meiosis terdiri atas meiosis I dan meiosis II serta menghasilkan variasi genetik.",
        ],
        "terms": [
            ("Kromatid", "Salah satu dari dua salinan kromosom yang terhubung di sentromer."),
            ("Pindah silang", "Pertukaran segmen kromosom homolog yang meningkatkan variasi."),
            ("Haploid", "Sel dengan satu set kromosom."),
        ],
        "case": "Kesalahan pemisahan kromosom pada meiosis disebut nondisjungsi. Nondisjungsi dapat menghasilkan gamet dengan jumlah kromosom abnormal.",
    },
    {
        "title": "DNA dan Replikasi",
        "summary": "DNA menyimpan informasi genetik dalam urutan basa. Replikasi DNA bersifat semikonservatif dan berlangsung sebelum pembelahan sel.",
        "points": [
            "DNA tersusun dari nukleotida yang mengandung basa A, T, G, dan C.",
            "Adenin berpasangan dengan timin, sedangkan guanin berpasangan dengan sitosin.",
            "Helikase membuka untai ganda DNA.",
            "DNA polimerase menambahkan nukleotida ke ujung 3'.",
            "Untai leading disintesis secara kontinu, sedangkan untai lagging disintesis dalam fragmen Okazaki.",
        ],
        "terms": [
            ("Semikonservatif", "Setiap molekul DNA baru mengandung satu untai lama dan satu untai baru."),
            ("Helikase", "Enzim yang membuka pilinan DNA."),
            ("Fragmen Okazaki", "Potongan pendek DNA pada sintesis untai lagging."),
        ],
        "case": "Enzim DNA polimerase memiliki aktivitas proofreading. Aktivitas ini menurunkan laju kesalahan replikasi, meskipun tidak menghilangkannya sepenuhnya.",
    },
    {
        "title": "Transkripsi dan Translasi",
        "summary": "Transkripsi menyalin DNA menjadi mRNA, sedangkan translasi menggunakan mRNA untuk merangkai asam amino menjadi protein. Proses ini menghubungkan gen dengan sifat fenotipik.",
        "points": [
            "RNA polimerase membuka DNA dan menambahkan nukleotida RNA.",
            "Pada eukariot, mRNA mengalami splicing, penambahan cap, dan poly-A tail.",
            "Kodon adalah tiga basa mRNA yang menentukan satu asam amino.",
            "tRNA membawa asam amino sesuai dengan antikodon.",
            "Ribosom memiliki situs A, P, dan E untuk translasi.",
        ],
        "terms": [
            ("Kodon", "Tiga urutan basa mRNA yang mengkode asam amino."),
            ("Intron", "Bagian mRNA yang dipotong sebelum translasi."),
            ("Antikodon", "Tiga basa pada tRNA yang berpasangan dengan kodon."),
        ],
        "case": "Kodon AUG berfungsi sebagai kodon start dan mengkode metionin pada banyak organisme. Kodon stop UAA, UAG, dan UGA tidak mengkode asam amino.",
    },
    {
        "title": "Hukum Mendel dan Pewarisan Sifat",
        "summary": "Hukum Mendel menjelaskan pewarisan sifat melalui segregasi alel dan pengelompokan bebas. Persilangan monohibrid dan dihibrid dapat digunakan untuk memprediksi perbandingan fenotipe.",
        "points": [
            "Alel dominan menutupi ekspresi alel resesif pada heterozigot.",
            "Hukum segregasi menyatakan bahwa dua alel berpisah saat pembentukan gamet.",
            "Hukum pengelompokan bebas berlaku untuk gen yang terletak pada kromosom berbeda.",
            "Persilangan monohibrid heterozigot menghasilkan perbandingan fenotipe 3:1.",
            "Persilangan dihibrid heterozigot menghasilkan perbandingan fenotipe 9:3:3:1.",
        ],
        "terms": [
            ("Genotipe", "Susunan alel yang dimiliki individu."),
            ("Fenotipe", "Sifat yang tampak atau teramati."),
            ("Heterozigot", "Individu dengan dua alel berbeda untuk satu gen."),
        ],
        "case": "Warna bunga pada tanaman tertentu dapat menunjukkan dominansi tidak penuh. Heterozigot menghasilkan fenotipe intermediet, bukan salah satu fenotipe homozigot.",
    },
    {
        "title": "Mutasi Genetik",
        "summary": "Mutasi adalah perubahan urutan DNA yang dapat terjadi secara spontan atau karena pengaruh mutagen. Mutasi dapat bersifat netral, menguntungkan, atau merugikan.",
        "points": [
            "Substitusi mengganti satu basa dengan basa lain.",
            "Insersi dan delesi dapat menyebabkan pergeseran kerangka baca.",
            "Mutasi titik dapat mengubah satu asam amino atau tidak mengubah asam amino sama sekali.",
            "Mutasi pada sel germinal dapat diwariskan, sedangkan mutasi somatik tidak diwariskan.",
            "Mutagen meliputi radiasi ultraviolet, bahan kimia, dan virus tertentu.",
        ],
        "terms": [
            ("Mutasi titik", "Perubahan pada satu pasangan basa DNA."),
            ("Frameshift", "Pergeseran kerangka baca akibat insersi atau delesi."),
            ("Mutagen", "Agen yang meningkatkan laju mutasi."),
        ],
        "case": "Mutasi delesi satu basa pada gen dapat mengubah seluruh rangkaian kodon setelah titik mutasi. Akibatnya, protein yang dihasilkan sering tidak berfungsi.",
    },
    {
        "title": "Enzim dan Kinetika Reaksi",
        "summary": "Enzim adalah biokatalis yang mempercepat reaksi tanpa ikut terpakai. Enzim bekerja dengan menurunkan energi aktivasi dan memiliki situs aktif yang spesifik.",
        "points": [
            "Model lock and key dan induced fit menjelaskan interaksi enzim-substrat.",
            "Laju reaksi enzim dipengaruhi oleh konsentrasi substrat, pH, suhu, dan inhibitor.",
            "Inhibitor kompetitif bersaing dengan substrat pada situs aktif.",
            "Inhibitor nonkompetitif mengubah bentuk enzim dan mengurangi aktivitas.",
            "Koenzim dan kofaktor membantu kerja enzim.",
        ],
        "terms": [
            ("Situs aktif", "Bagian enzim tempat substrat berikatan."),
            ("Energi aktivasi", "Energi minimum yang diperlukan agar reaksi berlangsung."),
            ("Inhibitor kompetitif", "Molekul yang menempati situs aktif dan menghalangi substrat."),
        ],
        "case": "Pada suhu sangat tinggi, enzim dapat mengalami denaturasi. Perubahan struktur tiga dimensi ini menyebabkan situs aktif tidak lagi cocok dengan substrat.",
    },
    {
        "title": "Metabolisme Karbohidrat",
        "summary": "Metabolisme karbohidrat mencakup pemecahan glukosa, penyimpanan glikogen, dan jalur alternatif seperti glukoneogenesis. Hormon insulin dan glukagon mengatur kadar glukosa darah.",
        "points": [
            "Glikolisis menghasilkan piruvat, ATP, dan NADH.",
            "Glikogen disimpan terutama di hati dan otot.",
            "Glikogenolisis memecah glikogen menjadi glukosa.",
            "Gluconeogenesis membentuk glukosa dari laktat, asam amino, dan gliserol.",
            "Insulin menurunkan kadar glukosa darah, sedangkan glukagon meningkatkannya.",
        ],
        "terms": [
            ("Glikogenesis", "Proses pembentukan glikogen dari glukosa."),
            ("Glikogenolisis", "Proses pemecahan glikogen."),
            ("Gluconeogenesis", "Sintesis glukosa dari prekursor nonkarbohidrat."),
        ],
        "case": "Saat berpuasa, hati memecah glikogen untuk menjaga kadar glukosa darah. Jika simpanan glikogen menipis, glukoneogenesis menjadi jalur penting.",
    },
    {
        "title": "Metabolisme Lipid",
        "summary": "Lipid berfungsi sebagai sumber energi, komponen membran, dan molekul sinyal. Oksidasi asam lemak terjadi melalui beta-oksidasi di mitokondria.",
        "points": [
            "Asam lemak diaktifkan menjadi asil-KoA sebelum masuk mitokondria.",
            "Beta-oksidasi memecah asil-KoA menjadi asetil-KoA dan menghasilkan NADH serta FADH2.",
            "Asetil-KoA masuk siklus Krebs untuk menghasilkan energi.",
            "Badan keton dibentuk di hati saat asetil-KoA berlebih.",
            "Kolesterol merupakan komponen penting membran dan prekursor hormon steroid.",
        ],
        "terms": [
            ("Beta-oksidasi", "Jalur pemecahan asam lemak di mitokondria."),
            ("Badan keton", "Bahan bakar alternatif yang diproduksi hati saat kelaparan."),
            ("Lipoprotein", "Partikel yang mengangkut lipid dalam darah."),
        ],
        "case": "Pada diabetes tipe 1 yang tidak terkontrol, pemecahan lemak meningkat dan produksi badan keton dapat berlebihan. Kondisi ini dapat menyebabkan ketoasidosis.",
    },
    {
        "title": "Sinyal Sel dan Reseptor",
        "summary": "Sel berkomunikasi melalui molekul sinyal yang berikatan dengan reseptor. Transduksi sinyal mengubah sinyal ekstraseluler menjadi respons intraseluler.",
        "points": [
            "Reseptor permukaan sel meliputi reseptor terkait protein G, reseptor tirosin kinase, dan kanal ion.",
            "Reseptor intraseluler dapat mengikat hormon steroid yang larut lipid.",
            "Transduksi sinyal sering melibatkan kaskade fosforilasi.",
            "Pembawa pesan kedua seperti cAMP dan kalsium memperkuat sinyal.",
            "Umpan balik negatif membantu menjaga homeostasis.",
        ],
        "terms": [
            ("Ligan", "Molekul sinyal yang berikatan dengan reseptor."),
            ("Pembawa pesan kedua", "Molekul intraseluler yang meneruskan sinyal."),
            ("Umpan balik negatif", "Mekanisme yang meredam perubahan untuk menjaga kestabilan."),
        ],
        "case": "Epinefrin berikatan dengan reseptor beta-adrenergik dan meningkatkan cAMP. Peningkatan cAMP memicu pemecahan glikogen untuk menyediakan energi cepat.",
    },
    {
        "title": "Jaringan Epitel dan Jaringan Ikatan",
        "summary": "Jaringan epitel melapisi permukaan tubuh dan membentuk kelenjar, sedangkan jaringan ikat menghubungkan, mendukung, dan melindungi organ. Struktur jaringan berkaitan erat dengan fungsinya.",
        "points": [
            "Epitel simpleks tersusun satu lapis sel, sedangkan epitel berlapis memiliki beberapa lapis.",
            "Epitel silindris sering ditemukan pada saluran pencernaan.",
            "Jaringan ikat longgar mengandung serat kolagen dan elastin.",
            "Jaringan ikat padat ditemukan pada tendon dan ligamen.",
            "Tulang dan darah merupakan bentuk khusus jaringan ikat.",
        ],
        "terms": [
            ("Epitel", "Jaringan yang melapisi permukaan dan rongga tubuh."),
            ("Kolagen", "Protein serat yang memberikan kekuatan pada jaringan ikat."),
            ("Matriks ekstraseluler", "Material di luar sel yang mendukung jaringan."),
        ],
        "case": "Epitel usus halus memiliki mikrovili untuk memperluas permukaan penyerapan. Struktur ini meningkatkan efisiensi penyerapan nutrisi.",
    },
    {
        "title": "Sistem Imun Dasar",
        "summary": "Sistem imun melindungi tubuh melalui imunitas bawaan dan adaptif. Imunitas bawaan bekerja cepat dan tidak spesifik, sedangkan imunitas adaptif mengenali antigen secara spesifik.",
        "points": [
            "Makrofag dan neutrofil melakukan fagositosis patogen.",
            "Sel dendritik mempresentasikan antigen kepada sel T.",
            "Limfosit B menghasilkan antibodi.",
            "Limfosit T helper mengaktifkan respons imun, sedangkan sel T sitotoksik membunuh sel terinfeksi.",
            "Memori imun mempercepat respons pada paparan antigen berikutnya.",
        ],
        "terms": [
            ("Antigen", "Zat yang memicu respons imun."),
            ("Antibodi", "Protein yang mengikat antigen secara spesifik."),
            ("Memori imun", "Kemampuan sistem imun mengingat antigen dari paparan sebelumnya."),
        ],
        "case": "Vaksinasi memperkenalkan antigen yang tidak berbahaya agar sistem imun membentuk memori. Pada paparan alami berikutnya, respons imun dapat terjadi lebih cepat dan kuat.",
    },
    {
        "title": "Bioteknologi dan PCR",
        "summary": "Bioteknologi memanfaatkan organisme atau komponennya untuk menghasilkan produk dan layanan. PCR adalah teknik memperbanyak DNA secara in vitro melalui siklus suhu.",
        "points": [
            "PCR terdiri atas denaturasi, annealing primer, dan ekstensi.",
            "Taq DNA polimerase tahan panas dan digunakan pada PCR.",
            "Gel elektroforesis memisahkan fragmen DNA berdasarkan ukuran.",
            "Kloning gen memasukkan gen target ke vektor plasmid.",
            "CRISPR-Cas9 memungkinkan penyuntingan DNA secara terarah.",
        ],
        "terms": [
            ("Primer", "Urutan DNA pendek yang menentukan daerah target amplifikasi."),
            ("Plasmid", "DNA sirkular yang digunakan sebagai vektor kloning."),
            ("CRISPR-Cas9", "Sistem penyuntingan gen yang dipandu RNA."),
        ],
        "case": " PCR dapat digunakan untuk mendeteksi materi genetik patogen dari sampel klinis. Kecepatan dan sensitivitasnya menjadikannya alat penting dalam diagnostik.",
    },
    {
        "title": "Evolusi dan Seleksi Alam",
        "summary": "Evolusi adalah perubahan frekuensi alel dalam populasi dari generasi ke generasi. Seleksi alam, mutasi, migrasi, dan hanyutan genetik merupakan mekanisme perubahan evolusi.",
        "points": [
            "Seleksi alam meningkatkan frekuensi sifat yang meningkatkan keberhasilan reproduksi.",
            "Mutasi menghasilkan variasi genetik baru.",
            "Hanyutan genetik lebih kuat pada populasi kecil.",
            "Migrasi dapat membawa alel baru ke populasi.",
            "Spesiasi terjadi ketika populasi terisolasi secara reproduktif.",
        ],
        "terms": [
            ("Frekuensi alel", "Proporsi alel tertentu dalam populasi."),
            ("Hanyutan genetik", "Perubahan frekuensi alel akibat peristiwa acak."),
            ("Spesiasi", "Pembentukan spesies baru dari populasi leluhur."),
        ],
        "case": "Resistensi antibiotik pada bakteri merupakan contoh seleksi alam yang cepat. Bakteri dengan sifat resisten lebih mampu bertahan dan bereproduksi saat antibiotik digunakan.",
    },
    {
        "title": "Ekologi Populasi dan Interaksi Spesies",
        "summary": "Ekologi populasi mempelajari ukuran, distribusi, dan pertumbuhan populasi. Interaksi antar spesies meliputi kompetisi, predasi, parasitisme, dan mutualisme.",
        "points": [
            "Pertumbuhan eksponensial terjadi saat sumber daya tidak terbatas.",
            "Pertumbuhan logistik dibatasi oleh daya dukung lingkungan.",
            "Kompetisi intraspesifik terjadi antarindividu spesies yang sama.",
            "Predasi dapat mengatur ukuran populasi mangsa dan predator.",
            "Mutualisme memberikan manfaat bagi kedua spesies yang berinteraksi.",
        ],
        "terms": [
            ("Daya dukung", "Jumlah maksimum individu yang dapat didukung lingkungan."),
            ("Relung ekologi", "Peran dan posisi spesies dalam ekosistem."),
            ("Mutualisme", "Interaksi yang menguntungkan kedua pihak."),
        ],
        "case": "Populasi yang terlalu padat dapat meningkatkan persaingan sumber daya dan penyebaran penyakit. Keduanya dapat menurunkan laju pertumbuhan populasi.",
    },
]


def build_text(topic):
    lines = [topic["title"], "", topic["summary"], "", "Poin penting:"]
    lines += [f"- {point}" for point in topic["points"]]
    lines += ["", "Istilah kunci:"]
    lines += [f"- {term}: {definition}" for term, definition in topic["terms"]]
    lines += ["", "Studi kasus:", topic["case"]]
    return "\n".join(lines)


def safe(slug):
    return "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in slug)


def write_pdf(path: Path, text: str):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    y = 64.0
    for raw_line in text.splitlines():
        lines = textwrap.wrap(raw_line, width=92) or [""]
        for line in lines:
            if y > 790:
                page = doc.new_page(width=595, height=842)
                y = 64.0
            page.insert_text((56, y), line, fontsize=10, fontname="helv")
            y += 14
    doc.save(path)
    doc.close()


def write_docx(path: Path, topic: dict):
    document = docx.Document()
    document.add_heading(topic["title"], level=1)
    document.add_paragraph(topic["summary"])
    document.add_heading("Poin penting", level=2)
    for point in topic["points"]:
        document.add_paragraph(point, style="List Bullet")
    document.add_heading("Istilah kunci", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Istilah"
    table.rows[0].cells[1].text = "Penjelasan"
    for term, definition in topic["terms"]:
        cells = table.add_row().cells
        cells[0].text = term
        cells[1].text = definition
    document.add_heading("Studi kasus", level=2)
    document.add_paragraph(topic["case"])
    document.save(path)


def write_xlsx(path: Path, topic: dict):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ringkasan"
    sheet.append(["Topik", topic["title"]])
    sheet.append(["Ringkasan", topic["summary"]])
    sheet.append([])
    sheet.append(["Poin penting", "Detail"])
    for point in topic["points"]:
        sheet.append(["Poin", point])
    sheet.append([])
    sheet.append(["Istilah", "Penjelasan"])
    for term, definition in topic["terms"]:
        sheet.append([term, definition])
    sheet.append([])
    sheet.append(["Studi kasus", topic["case"]])
    for column, width in {"A": 28, "B": 100}.items():
        sheet.column_dimensions[column].width = width
    workbook.save(path)


def write_csv(path: Path, topic: dict):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Bagian", "Isi"])
        writer.writerow(["Topik", topic["title"]])
        writer.writerow(["Ringkasan", topic["summary"]])
        for point in topic["points"]:
            writer.writerow(["Poin penting", point])
        for term, definition in topic["terms"]:
            writer.writerow(["Istilah", f"{term}: {definition}"])
        writer.writerow(["Studi kasus", topic["case"]])


def write_pptx(path: Path, topic: dict):
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = topic["title"]
    title_slide.placeholders[1].text = topic["summary"]

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Poin Penting"
    body = slide.placeholders[1].text_frame
    body.text = topic["points"][0]
    for point in topic["points"][1:]:
        paragraph = body.add_paragraph()
        paragraph.text = point
        paragraph.level = 0

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Istilah Kunci"
    body = slide.placeholders[1].text_frame
    body.text = f"{topic['terms'][0][0]}: {topic['terms'][0][1]}"
    for term, definition in topic["terms"][1:]:
        paragraph = body.add_paragraph()
        paragraph.text = f"{term}: {definition}"

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Studi Kasus"
    slide.placeholders[1].text = topic["case"]
    presentation.save(path)


FORMAT_PLAN = [
    "pdf", "pdf", "pdf", "pdf",
    "docx", "docx", "docx",
    "xlsx", "xlsx", "xlsx",
    "txt", "txt", "txt",
    "md", "md", "md",
    "csv", "csv",
    "pptx", "pptx",
]


def generate_files():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generated = []
    for index, (topic, fmt) in enumerate(zip(TOPICS, FORMAT_PLAN), start=1):
        name = f"{index:02d}_{safe(topic['title']).strip().replace(' ', '_')}.{fmt}"
        path = OUT_DIR / name
        text = build_text(topic)
        if fmt == "pdf":
            write_pdf(path, text)
        elif fmt == "docx":
            write_docx(path, topic)
        elif fmt == "xlsx":
            write_xlsx(path, topic)
        elif fmt == "csv":
            write_csv(path, topic)
        elif fmt == "pptx":
            write_pptx(path, topic)
        else:
            path.write_text(text, encoding="utf-8")
        generated.append((path, fmt, topic["title"]))
    return generated


def api_json(method, path, **kwargs):
    response = requests.request(method, f"{API}{path}", timeout=60, **kwargs)
    try:
        data = response.json()
    except Exception:
        data = {"raw": response.text}
    if response.status_code >= 400:
        raise RuntimeError(f"{method} {path} -> {response.status_code}: {data}")
    return data, response


def ensure_subject():
    data, _ = api_json("GET", "/subjects")
    for subject in data.get("subjects", []):
        if subject.get("name") == SUBJECT_NAME:
            return subject["id"]
    api_json("POST", "/subjects", data={"name": SUBJECT_NAME, **SUBJECT_META})
    data, _ = api_json("GET", "/subjects")
    for subject in data.get("subjects", []):
        if subject.get("name") == SUBJECT_NAME:
            return subject["id"]
    raise RuntimeError("Subject demo tidak ditemukan setelah dibuat")


def main():
    files = generate_files()
    subject_id = ensure_subject()
    print(f"Subject demo id={subject_id}, {len(files)} file dibuat di {OUT_DIR}")

    current, _ = api_json("GET", f"/materials?subject_id={subject_id}&limit=200")
    existing = {item.get("filename") for item in current.get("materials", [])}
    for path, fmt, title in files:
        if path.name in existing:
            print(f"SKIP  {path.name}")
            continue
        with path.open("rb") as handle:
            response = requests.post(
                f"{API}/materials",
                data={"subject_id": str(subject_id)},
                files={"file": (path.name, handle, "application/octet-stream")},
                timeout=120,
            )
        try:
            payload = response.json()
        except Exception:
            payload = response.text
        print(f"UPLOAD {path.name} [{fmt}] -> {response.status_code} {payload}")
        if response.status_code >= 400:
            raise RuntimeError(f"Upload gagal: {path.name}")

    print("\nMenunggu worker memproses materi...")
    deadline = time.time() + 15 * 60
    last_report = 0
    while time.time() < deadline:
        data, _ = api_json("GET", f"/materials?subject_id={subject_id}&limit=200")
        materials = data.get("materials", [])
        statuses = {}
        for item in materials:
            statuses[item.get("status")] = statuses.get(item.get("status"), 0) + 1
        if time.time() - last_report > 10:
            print(f"  status: {statuses} (total {len(materials)})")
            last_report = time.time()
        if len(materials) >= len(files) and all(
            item.get("status") in {"Ready", "Failed", "Trashed"} for item in materials
        ):
            break
        time.sleep(3)

    data, _ = api_json("GET", f"/materials?subject_id={subject_id}&limit=200")
    materials = data.get("materials", [])
    print("\nRingkasan materi:")
    for item in sorted(materials, key=lambda x: x.get("id", 0)):
        print(f"  #{item.get('id'):>2} {item.get('status'):<10} {item.get('material_type'):<8} {item.get('filename')}")
    failed = [item for item in materials if item.get("status") != "Ready"]
    if failed:
        print(f"\nPERINGATAN: {len(failed)} materi tidak Ready")
        for item in failed:
            print(" ", item.get("id"), item.get("status"), item.get("failure_reason"), item.get("filename"))
        return 1
    print(f"\nOK: {len(materials)} materi Ready untuk subject {subject_id}")
    print(json.dumps({"subject_id": subject_id, "materials": len(materials)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
